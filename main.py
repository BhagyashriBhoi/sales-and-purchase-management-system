import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

file_name = "sales_and_purchase.xlsx"

# Create the file if it doesn't exist
if not os.path.exists(file_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales and Purchase"
    ws.append([
        "Product Name", "Purchase Quantity", "Purchase Price", "Total",
        "Sales Quantity", "Selling Price", "Stock Available"
    ])
    wb.save(file_name)
    wb.close()

root = tk.Tk()
root.title("Login Page")
root.config(bg="lightblue")
root.geometry("800x600")

def purchase_product():
    add_window = tk.Toplevel()
    add_window.title("Purchase Product")
    add_window.geometry("500x400")

    tk.Label(add_window, text="Product Name").pack()
    pname = tk.Entry(add_window)
    pname.pack()

    tk.Label(add_window, text="Purchase Quantity").pack()
    qty = tk.Entry(add_window)
    qty.pack()

    tk.Label(add_window, text="Purchase Price").pack()
    price = tk.Entry(add_window)
    price.pack()

    def save_purchase():
        wb = load_workbook(file_name)
        ws = wb.active
        product = pname.get()
        
        try:
            quantity = int(qty.get())
            purchase_price = float(price.get())
        except ValueError:
            messagebox.showerror("Error", "Please enter valid numbers.")
            return

        found = False

        for row in ws.iter_rows(min_row=2):
            if row[0].value == product:
                old_qty = row[1].value or 0
                stock = row[6].value or 0

                row[1].value = old_qty + quantity
                row[2].value = purchase_price
                row[3].value = (old_qty + quantity) * purchase_price
                row[6].value = stock + quantity

                found = True
                break

        if not found:
            ws.append([
                product, quantity, purchase_price, quantity * purchase_price, 0, 0, quantity
            ])

        wb.save(file_name)
        wb.close()
        messagebox.showinfo("Success", "Product Added")
        add_window.destroy()

    tk.Button(add_window,text="purchase",command=save_purchase).pack(pady=10)
    tk.Button(add_window, text="Back", command=add_window.destroy, bg="lightgrey", width=15).pack(pady=5)

def sale_product():
    sale_window = tk.Toplevel()
    sale_window.title("Sale Product")
    sale_window.geometry("500x400")

    tk.Label(sale_window, text="Product Name").pack()
    pname = tk.Entry(sale_window)
    pname.pack()

    tk.Label(sale_window, text="Sale Quantity").pack()
    qty = tk.Entry(sale_window)
    qty.pack()

    tk.Label(sale_window, text="Sale Price").pack()
    price = tk.Entry(sale_window)
    price.pack()

    def save_sale():
        wb = load_workbook(file_name)
        ws = wb.active

        product = pname.get()
        try:
            quantity = int(qty.get())
            sale_price = float(price.get())
        except ValueError:
            messagebox.showerror("Error", "Please enter valid numbers.")
            return

        for row in ws.iter_rows(min_row=2):
            if row[0].value == product:
                stock = row[6].value or 0

                if stock >= quantity:
                    current_sales_qty = row[4].value or 0
                    row[4].value = current_sales_qty + quantity
                    row[5].value = sale_price
                    row[6].value = stock - quantity

                    wb.save(file_name)
                    wb.close()

                    messagebox.showinfo("Success", "Sale Completed")
                    sale_window.destroy()
                    return
                else:
                    messagebox.showerror("Error", "Not Enough Stock")
                    return

        messagebox.showerror("Error", "Product Not Found")

    tk.Button(sale_window, text="Sale", command=save_sale).pack(pady=10)
    tk.Button(sale_window, text="Back", command=sale_window.destroy, bg="lightgrey", width=15).pack(pady=5)


def view_stock():
    wb = load_workbook(file_name)
    ws = wb.active

    stock_window = tk.Toplevel()
    stock_window.title("View Stock")
    stock_window.geometry("400x500")

    text = tk.Text(stock_window)
    text.pack(fill="both", expand=True)

    text.insert("end", "Product\t\tStock\n")
    text.insert("end", "-"*40+"\n")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            text.insert("end", f"{row[0]}\t\t{row[6]}\n")

    wb.close()
    tk.Button(stock_window, text="Back", command=stock_window.destroy, bg="lightgrey", width=15).pack(pady=5)


def show_all():
    wb = load_workbook(file_name)
    ws = wb.active

    window = tk.Toplevel()
    window.title("All Products")
    window.geometry("1000x500")

    frame = tk.Frame(window)
    frame.pack(fill="both", expand=True)

    columns = ("Product", "Purchase Qty", "Purchase Price", "Total", "Sale Qty", "Sale Price", "Stock")

    tree = ttk.Treeview(frame, columns=columns, show="headings")
    vs = tk.Scrollbar(frame, orient="vertical", command=tree.yview)
    hs = tk.Scrollbar(frame, orient="horizontal", command=tree.xview)
    tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=130, anchor="center")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            tree.insert("", tk.END, values=row)

    tree.grid(row=0, column=0, sticky="nsew")
    vs.grid(row=0, column=1, sticky="ns")
    hs.grid(row=1, column=0, sticky="ew")

    frame.rowconfigure(0, weight=1)
    frame.columnconfigure(0, weight=1)
    wb.close()
    tk.Button(window, text="Back", command=window.destroy, bg="lightgrey", width=15).pack(pady=5)


# NEW FEATURE: UPDATE PRODUCT
def update_product():
    upd_window = tk.Toplevel()
    upd_window.title("Update Product")
    upd_window.geometry("500x400")

    tk.Label(upd_window, text="Product Name (Exact Match):").pack(pady=5)
    pname = tk.Entry(upd_window)
    pname.pack()

    tk.Label(upd_window, text="New Purchase Price (Leave blank to skip):").pack(pady=5)
    pprice = tk.Entry(upd_window)
    pprice.pack()

    tk.Label(upd_window, text="New Sale Price (Leave blank to skip):").pack(pady=5)
    sprice = tk.Entry(upd_window)
    sprice.pack()

    tk.Label(upd_window, text="Fix/Update Stock (Leave blank to skip):").pack(pady=5)
    nstock = tk.Entry(upd_window)
    nstock.pack()

    def save_update():
        product = pname.get()
        if not product:
            messagebox.showerror("Error", "Product name is required")
            return

        wb = load_workbook(file_name)
        ws = wb.active
        found = False

        for row in ws.iter_rows(min_row=2):
            if row[0].value == product:
                found = True
                
                if pprice.get():
                    try:
                        row[2].value = float(pprice.get())
                        qty = row[1].value or 0
                        row[3].value = qty * row[2].value # update total cost
                    except ValueError:
                        messagebox.showerror("Error", "Invalid Purchase Price")
                        wb.close()
                        return
                
                if sprice.get():
                    try:
                        row[5].value = float(sprice.get())
                    except ValueError:
                        messagebox.showerror("Error", "Invalid Sale Price")
                        wb.close()
                        return

                if nstock.get():
                    try:
                        row[6].value = int(nstock.get())
                    except ValueError:
                        messagebox.showerror("Error", "Invalid Stock")
                        wb.close()
                        return

                wb.save(file_name)
                messagebox.showinfo("Success", f"{product} updated successfully!")
                upd_window.destroy()
                break

        if not found:
            messagebox.showerror("Error", "Product not found")
        wb.close()

    tk.Button(upd_window, text="Update Data", command=save_update, bg="orange").pack(pady=20)
    tk.Button(upd_window, text="Back", command=upd_window.destroy, bg="lightgrey", width=15).pack(pady=5)


# NEW FEATURE: DELETE PRODUCT
def delete_product():
    del_window = tk.Toplevel()
    del_window.title("Delete Product")
    del_window.geometry("400x200")

    tk.Label(del_window, text="Enter Product Name to Delete:").pack(pady=10)
    pname = tk.Entry(del_window)
    pname.pack(pady=5)

    def perform_delete():
        product = pname.get()
        if not product:
            return

        wb = load_workbook(file_name)
        ws = wb.active
        found = False

        # Enumerate to track the exact row number for deletion
        for index, row in enumerate(ws.iter_rows(min_row=2), start=2):
            if row[0].value == product:
                ws.delete_rows(index)
                wb.save(file_name)
                messagebox.showinfo("Success", f"{product} deleted completely.")
                del_window.destroy()
                found = True
                break

        if not found:
            messagebox.showerror("Error", "Product not found")
        wb.close()

    tk.Button(del_window, text="Delete Product", command=perform_delete, bg="red", fg="white").pack(pady=20)
    tk.Button(del_window, text="Back", command=del_window.destroy, bg="lightgrey", width=15).pack(pady=5)


def Dashboard():
    user = uname.get()
    password = passe.get()
    
    if user == "admin" and password == "1234":
        root.withdraw()
        window = tk.Toplevel()
        window.title("Dashboard")
        window.geometry("800x800")
        window.config(bg="lightblue")

        def logout():
            window.destroy()          # Close dashboard
            uname.delete(0, tk.END)   # Clear username field
            passe.delete(0, tk.END)   # Clear password field
            root.deiconify()
        
        dash = tk.Label(window, text="Welcome to Dashboard", font=("Times New Roman", 40), bg="lightblue")
        dash.pack(pady=20)
        
        tk.Button(window, text="Purchase Product", font=("Times New Roman", 20), bg="grey", fg="white", width=20, command=purchase_product).pack(pady=10)
        tk.Button(window, text="Sale Product", font=("Times New Roman", 20), bg="grey", fg="white", width=20, command=sale_product).pack(pady=10)
        tk.Button(window, text="Update Product", font=("Times New Roman", 20), bg="orange", width=20, command=update_product).pack(pady=10)
        tk.Button(window, text="View Product Stock", font=("Times New Roman", 20), bg="grey", fg="white", width=20, command=view_stock).pack(pady=10)
        tk.Button(window, text="View All Products", font=("Times New Roman", 20), bg="grey", fg="white", width=20, command=show_all).pack(pady=10)
        tk.Button(window, text="Delete Product", font=("Times New Roman", 20), bg="red", fg="white", width=20, command=delete_product).pack(pady=10)
        tk.Button(window, text="Back to Login", font=("Times New Roman", 20), bg="darkblue", fg="white", width=20, command=logout).pack(pady=20)
    else:
        messagebox.showerror("Login Failed", "Invalid username or password")


# --- Login UI ---
login = tk.Label(root, text="Login", font=("Times New Roman", 40), bg="lightblue")
login.pack(pady=20)

username = tk.Label(root, text="Enter your username", font=("Times New Roman", 20), bg="lightblue")
username.pack()
uname = tk.Entry(root, font=("Times New Roman", 20))
uname.pack(pady=10)

password_lbl = tk.Label(root, text="Enter your password", font=("Times New Roman", 20), bg="lightblue")
password_lbl.pack()
passe = tk.Entry(root, font=("Times New Roman", 20), show="*")
passe.pack(pady=10)

b1 = tk.Button(root, text="Login", font=("Times New Roman", 22), bg="light green", command=Dashboard)
b1.pack(pady=20)

root.mainloop()